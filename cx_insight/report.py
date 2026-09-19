"""Excel report writer.

This is the one export, not the product — the source of truth is the DuckDB warehouse.
Three sheets: Executive Dashboard, Priority Tickets, and Full Ticket Log.
"""
from __future__ import annotations

import os
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PRIORITY_COLORS = {"Critical": "C00000", "High": "E26B0A", "Medium": "F4B942", "Low": "70AD47"}
SENTIMENT_COLORS = {"Positive": "70AD47", "Neutral": "9DC3E6", "Negative": "E26B0A", "Urgent": "C00000"}


def _fill(hexc):
    return PatternFill(start_color=hexc, end_color=hexc, fill_type="solid")


def _border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def write_report(df: pd.DataFrame, kpis: dict, exec_summary: str, output_path) -> None:
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = Workbook()

    _dashboard(wb.active, df, kpis, exec_summary)
    _priority_sheet(wb.create_sheet("Priority Tickets"), df)
    _full_log(wb.create_sheet("Full Ticket Log"), df)

    wb.save(output_path)


def _dashboard(ws, df, kpis, exec_summary):
    ws.title = "Executive Dashboard"
    ws.sheet_view.showGridLines = False
    for col in "ABCD":
        ws.column_dimensions[col].width = 29

    ws.merge_cells("A1:D1")
    ws["A1"] = "CX INSIGHT REPORT  —  AI-Powered Support Ticket Analysis"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = _fill("1F3864")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:D2")
    span = f"{df['date'].min().strftime('%b %d')} – {df['date'].max().strftime('%b %d, %Y')}"
    ws["A2"] = f"Generated {datetime.now():%B %d, %Y}  |  Analysis period {span}"
    ws["A2"].font = Font(size=10, color="FFFFFF")
    ws["A2"].fill = _fill("2E4D8A")
    ws["A2"].alignment = Alignment(horizontal="center")

    _band(ws, 4, "KEY PERFORMANCE INDICATORS")
    row = 5
    items = list(kpis.items())
    for i in range(0, len(items), 2):
        ws.row_dimensions[row].height = 38
        for j, (k, v) in enumerate(items[i:i + 2]):
            a, b = get_column_letter(1 + j * 2), get_column_letter(2 + j * 2)
            ws.merge_cells(f"{a}{row}:{b}{row}")
            cell = ws[f"{a}{row}"]
            cell.value = f"{k}:  {v}"
            cell.font = Font(bold=True, size=11, color="FFFFFF")
            cell.fill = _fill("1F3864" if (i + j) % 2 == 0 else "2E4D8A")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = _border()
        row += 1

    row += 1
    _band(ws, row, "EXECUTIVE SUMMARY")
    row += 1
    ws.merge_cells(f"A{row}:D{row + 4}")
    ws[f"A{row}"] = exec_summary
    ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws[f"A{row}"].font = Font(size=10)
    ws[f"A{row}"].fill = _fill("EBF3FB")
    row += 6

    _band(ws, row, "TICKET VOLUME BY CATEGORY")
    row += 1
    for h, col in [("Category", "A"), ("Count", "B"), ("% of Total", "C"), ("Top Priority", "D")]:
        c = ws[f"{col}{row}"]
        c.value = h
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = _fill("1F3864")
        c.alignment = Alignment(horizontal="center")
        c.border = _border()
    row += 1
    summ = df.groupby("ai_category").agg(
        count=("ticket_id", "count"),
        top=("ai_priority", lambda x: x.value_counts().idxmax()),
    ).sort_values("count", ascending=False)
    for i, (cat, r) in enumerate(summ.iterrows()):
        shade = "EBF3FB" if i % 2 == 0 else "FFFFFF"
        for col, val, align in [("A", cat, "left"), ("B", int(r["count"]), "center"),
                                 ("C", f"{round(r['count'] / len(df) * 100, 1)}%", "center"),
                                 ("D", r["top"], "center")]:
            c = ws[f"{col}{row}"]
            c.value = val
            c.fill = _fill(shade)
            c.font = Font(size=10)
            c.alignment = Alignment(horizontal=align)
            c.border = _border()
        row += 1


def _band(ws, row, text):
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"] = text
    ws[f"A{row}"].font = Font(bold=True, size=11, color="FFFFFF")
    ws[f"A{row}"].fill = _fill("2E75B6")
    ws[f"A{row}"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[row].height = 20


_COLS = ["Ticket ID", "Date", "Segment", "Channel", "Subject", "Category", "Priority", "AI Summary", "Business Impact"]


def _priority_sheet(ws, df):
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHI", [10, 12, 14, 10, 22, 22, 12, 40, 34]):
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:I1")
    ws["A1"] = "PRIORITY TICKETS — Critical & High"
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = _fill("C00000")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    _header_row(ws, _COLS, 2, "1F3864")

    pri = df[df["ai_priority"].isin(["Critical", "High"])].sort_values(
        "ai_priority", key=lambda x: x.map({"Critical": 0, "High": 1})
    )
    for i, (_, r) in enumerate(pri.iterrows()):
        rn = i + 3
        vals = [r["ticket_id"], r["date"].strftime("%Y-%m-%d"), r["customer_segment"], r["channel"],
                r["subject"], r["ai_category"], r["ai_priority"], r["ai_summary"], r["ai_business_impact"]]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=rn, column=ci, value=v)
            c.border = _border()
            c.font = Font(size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left" if ci in (5, 8, 9) else "center")
            if ci == 7:
                c.font = Font(bold=True, size=9, color="FFFFFF")
                c.fill = _fill(PRIORITY_COLORS.get(r["ai_priority"], "000000"))
            elif i % 2 == 0:
                c.fill = _fill("FFF2CC" if r["ai_priority"] == "Critical" else "FCE4D6")
        ws.row_dimensions[rn].height = 38


def _full_log(ws, df):
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHI", [10, 12, 14, 10, 24, 22, 12, 13, 40]):
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:I1")
    ws["A1"] = "COMPLETE TICKET LOG — All Tickets with AI Classification"
    ws["A1"].font = Font(bold=True, size=12, color="FFFFFF")
    ws["A1"].fill = _fill("1F3864")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    cols = _COLS[:7] + ["Sentiment", "AI Summary"]
    _header_row(ws, cols, 2, "2E75B6")

    for i, (_, r) in enumerate(df.iterrows()):
        rn = i + 3
        vals = [r["ticket_id"], r["date"].strftime("%Y-%m-%d"), r["customer_segment"], r["channel"],
                r["subject"], r["ai_category"], r["ai_priority"], r["ai_sentiment"], r["ai_summary"]]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=rn, column=ci, value=v)
            c.border = _border()
            c.font = Font(size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left" if ci in (5, 9) else "center")
            if ci == 7:
                c.font = Font(bold=True, size=9, color="FFFFFF")
                c.fill = _fill(PRIORITY_COLORS.get(r["ai_priority"], "CCCCCC"))
            elif ci == 8:
                c.font = Font(bold=True, size=9, color="FFFFFF")
                c.fill = _fill(SENTIMENT_COLORS.get(r["ai_sentiment"], "CCCCCC"))
            else:
                c.fill = _fill("EBF3FB" if i % 2 == 0 else "FFFFFF")
        ws.row_dimensions[rn].height = 28


def _header_row(ws, cols, row, color):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = _fill(color)
        c.alignment = Alignment(horizontal="center")
        c.border = _border()
    ws.row_dimensions[row].height = 18
